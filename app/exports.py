"""Export saved records as spreadsheets, with formula injection protection."""
import csv
import io
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

COLUMNS = (("title", "Title"), ("agency", "Agency"), ("solicitation", "Solicitation"), ("naics", "NAICS"),
           ("set_aside", "Set-aside"), ("location", "Location"), ("posted", "Posted"), ("deadline", "Response deadline"),
           ("type", "Notice type"), ("status", "Review status"), ("notes", "Notes"), ("url", "SAM.gov URL"),
           ("saved_at", "Saved at (UTC)"), ("fetched_at", "Source snapshot (UTC)"),
           ("match_score", "Relevance score (not eligibility)"), ("match_evidence", "Score explanation"))


def safe_text(value):
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", "" if value is None else str(value))
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")) else value


def export_csv(items):
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([label for _, label in COLUMNS])
    writer.writerows([[safe_text(item.get(key)) for key, _ in COLUMNS] for item in items])
    return io.BytesIO(buffer.getvalue().encode("utf-8-sig"))


def export_excel(items):
    book = Workbook()
    sheet = book.active
    sheet.title = "Saved opportunities"
    sheet.append([label for _, label in COLUMNS])
    for item in items:
        sheet.append([safe_text(item.get(key))[:32767] for key, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="102D43")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 25
    sheet.column_dimensions["A"].width = 55
    sheet.column_dimensions["K"].width = 55
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer
